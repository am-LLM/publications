#!/usr/bin/env python3
from __future__ import annotations

import html
from pathlib import Path
from xml.sax.saxutils import escape

import markdown
from lxml import html as lhtml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
PAPERS = ROOT / 'papers'
FIGURES = PAPERS / 'figures'
OUT = PAPERS / 'submission'


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar'); fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    fld_end = OxmlElement('w:fldChar'); fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin); run._r.append(instr); run._r.append(fld_end)


def setup_document(title):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.68)
    sec.left_margin = Inches(.82); sec.right_margin = Inches(.82)
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'DejaVu Serif'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'DejaVu Serif'); normal.font.size = Pt(10.2)
    normal.paragraph_format.line_spacing = 1.15; normal.paragraph_format.space_after = Pt(5)
    for name, font, size, color in [('Title','DejaVu Serif',19,'17324D'),('Subtitle','DejaVu Serif',11,'52606D'),('Heading 1','DejaVu Sans',13,'17324D'),('Heading 2','DejaVu Sans',11,'234E70'),('Heading 3','DejaVu Sans',10,'52606D')]:
        st = styles[name]; st.font.name=font; st._element.rPr.rFonts.set(qn('w:eastAsia'),font); st.font.size=Pt(size); st.font.bold=name!='Subtitle'; st.font.color.rgb=RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(12 if name=='Heading 1' else 8); st.paragraph_format.space_after=Pt(5); st.paragraph_format.keep_with_next=True
    if 'Caption' not in [s.name for s in styles]: styles.add_style('Caption', WD_STYLE_TYPE.PARAGRAPH)
    cap=styles['Caption']; cap.font.name='DejaVu Serif'; cap._element.rPr.rFonts.set(qn('w:eastAsia'),'DejaVu Serif'); cap.font.size=Pt(8.5); cap.font.italic=True; cap.font.color.rgb=RGBColor(82,96,109); cap.paragraph_format.space_before=Pt(3); cap.paragraph_format.space_after=Pt(8)
    if 'Code' not in [s.name for s in styles]: styles.add_style('Code', WD_STYLE_TYPE.PARAGRAPH)
    code=styles['Code']; code.font.name='DejaVu Sans Mono'; code._element.rPr.rFonts.set(qn('w:eastAsia'),'DejaVu Sans Mono'); code.font.size=Pt(8); code.paragraph_format.left_indent=Inches(.2); code.paragraph_format.space_after=Pt(6)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT; footer.style='Caption'; footer.add_run(f'{title[:55]}  |  Page '); add_page_field(footer)
    doc.core_properties.title=title; doc.core_properties.author='Ali Imran Malik'; doc.core_properties.subject='Tiny Wonder / LoRA Lens Lab publication revision'
    return doc


def add_inline(paragraph, node):
    if node.text: paragraph.add_run(node.text)
    for child in node:
        tag = child.tag.lower() if isinstance(child.tag,str) else ''
        run = paragraph.add_run(''.join(child.itertext()))
        if tag in ('strong','b'): run.bold=True
        if tag in ('em','i'): run.italic=True
        if tag in ('code','tt'): run.font.name='DejaVu Sans Mono'; run.font.size=Pt(8.5)
        if tag == 'a': run.underline=True
        if child.tail: paragraph.add_run(child.tail)


def text_of(node): return ' '.join(''.join(node.itertext()).split())


def add_table(doc, node):
    rows=[]
    for tr in node.xpath('.//tr'):
        rows.append([text_of(c) for c in tr.xpath('./th|./td')])
    if not rows: return
    table=doc.add_table(rows=len(rows), cols=max(len(r) for r in rows)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
    for i,row in enumerate(rows):
        for j,val in enumerate(row):
            cell=table.cell(i,j); cell.text=val; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.0
                for r in p.runs: r.font.name='DejaVu Serif'; r.font.size=Pt(8); r.bold=(i==0)
            if i==0: set_cell_shading(cell,'D9EAF7')
    doc.add_paragraph().paragraph_format.space_after=Pt(2)


def add_image(doc, img_node):
    src = Path(img_node.get('src',''))
    png = FIGURES / 'rendered' / (src.stem + '.png')
    if png.exists():
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(png), width=Inches(6.15))
        p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(2)
    else:
        doc.add_paragraph(f'[Figure unavailable: {src}]', style='Caption')


def build(md_path, out_path):
    src=md_path.read_text(encoding='utf-8')
    rendered=markdown.markdown(src, extensions=['tables','fenced_code'])
    root=lhtml.fromstring(f'<div>{rendered}</div>')
    blocks=list(root); title_node=next((n for n in blocks if n.tag=='h1'),None); title=text_of(title_node) if title_node is not None else md_path.stem
    h2s=[n for n in blocks if n.tag=='h2']; subtitle_node=h2s[0] if h2s and text_of(h2s[0]).lower()!='abstract' else None
    abstract_node=next((n for n in blocks if n.tag in ('h2','h3') and text_of(n).lower()=='abstract'),None)
    doc=setup_document(title)
    p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(title)
    if subtitle_node is not None:
        p=doc.add_paragraph(style='Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(text_of(subtitle_node))
    for t in ('Ali Imran Malik','Independent Researcher, Islamabad, Pakistan','CEO, Hydrabotics Group','GitHub: github.com/am-llm · LinkedIn: linkedin.com/in/aliimranmalik'):
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.style='Caption'; p.add_run(t)
    if abstract_node is not None: doc.add_heading('Abstract', level=1)
    skipped_title=skipped_subtitle=skipped_abstract=False
    for node in blocks:
        tag=node.tag.lower() if isinstance(node.tag,str) else ''
        if tag=='h1' and not skipped_title: skipped_title=True; continue
        if subtitle_node is not None and node is subtitle_node and not skipped_subtitle: skipped_subtitle=True; continue
        if abstract_node is not None and node is abstract_node and not skipped_abstract: skipped_abstract=True; continue
        if tag in ('h2','h3','h4'):
            doc.add_heading(text_of(node), level={'h2':1,'h3':2,'h4':3}[tag]); continue
        if tag=='p':
            imgs=node.xpath('.//img')
            if imgs:
                add_image(doc, imgs[0]); continue
            p=doc.add_paragraph(style='Caption' if node.xpath('.//em') and len(text_of(node))<450 else 'Normal'); add_inline(p,node); continue
        if tag=='table': add_table(doc,node); continue
        if tag in ('ul','ol'):
            for li in node.xpath('./li'):
                p=doc.add_paragraph(style='List Bullet' if tag=='ul' else 'List Number'); add_inline(p,li)
            continue
        if tag=='pre': doc.add_paragraph(html.unescape(text_of(node)), style='Code'); continue
        if tag=='hr': doc.add_paragraph()
    doc.save(out_path)

if __name__=='__main__':
    build(PAPERS/'paper1_zero_base_host_draft.md', OUT/'paper1_zero_base_host.docx')
    build(PAPERS/'paper2_stacking_attenuation_draft.md', OUT/'paper2_stacking_attenuation.docx')
    print('built two docx files')
